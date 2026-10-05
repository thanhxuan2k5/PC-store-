import os
import re
from typing import Optional

def extract_text_from_file(file_path: str, filename: str) -> str:
    """
    Trích xuất toàn bộ văn bản từ file Word (.docx), PDF (.pdf), Markdown (.md, .markdown), và Text (.txt).
    """
    ext = os.path.splitext(filename)[1].lower()
    text = ""

    if ext == ".pdf":
        text = _extract_pdf(file_path)
    elif ext in [".docx", ".doc"]:
        text = _extract_docx(file_path)
    elif ext in [".md", ".markdown", ".txt"]:
        text = _extract_text_file(file_path)
    else:
        # Fallback thử đọc dạng text thông thường
        try:
            text = _extract_text_file(file_path)
        except Exception:
            raise ValueError(f"Định dạng file '{ext}' không được hỗ trợ. Vui lòng tải lên file .pdf, .docx, .md hoặc .txt")

    # Làm sạch văn bản: chuẩn hóa khoảng trắng, loại bỏ dòng trống thừa
    text = re.sub(r'\r\n', '\n', text)
    text = re.sub(r'\n{3,}', '\n\n', text)
    text = text.strip()

    if not text or len(text) < 10:
        raise ValueError("File tài liệu không chứa nội dung văn bản có thể trích xuất được hoặc file rỗng.")

    return text

def _extract_pdf(file_path: str) -> str:
    """Trích xuất text từ file PDF sử dụng pypdf"""
    import pypdf
    extracted_pages = []
    with open(file_path, "rb") as f:
        reader = pypdf.PdfReader(f)
        for page_num, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text and page_text.strip():
                extracted_pages.append(page_text.strip())
    
    return "\n\n".join(extracted_pages)

def _extract_docx(file_path: str) -> str:
    """Trích xuất text từ file Word .docx sử dụng python-docx"""
    import docx
    doc = docx.Document(file_path)
    paragraphs = []
    
    # 1. Đoạn văn thông thường
    for p in doc.paragraphs:
        p_text = p.text.strip()
        if p_text:
            paragraphs.append(p_text)
            
    # 2. Bảng biểu (Tables)
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
            if row_text:
                paragraphs.append(row_text)

    return "\n\n".join(paragraphs)

def _extract_text_file(file_path: str) -> str:
    """Trích xuất text từ file .md, .markdown, .txt với encoding fallback"""
    encodings = ["utf-8", "utf-8-sig", "utf-16", "cp1252", "latin-1"]
    for enc in encodings:
        try:
            with open(file_path, "r", encoding=enc) as f:
                return f.read()
        except (UnicodeDecodeError, Exception):
            continue
    
    raise ValueError("Không thể giải mã định dạng text của file.")
